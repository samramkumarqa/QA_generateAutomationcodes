"""Generate a project from the practice shop, npm-install it and run it with Playwright.

Slow (npm install + browser); skipped when node/npm aren't available.
Set PW_TEST_VERSION to pin @playwright/test to the browser build installed on this machine.
"""

import io
import os
import shutil
import subprocess

import pytest
from openpyxl import load_workbook

from fakes import FakePlannerLLM
from pwgen.planner import plan_tests
from pwgen.renderer import build_project

pytestmark = pytest.mark.skipif(not shutil.which("npm"), reason="npm not installed")


def _run(cmd, cwd, env=None, timeout=600):
    return subprocess.run(cmd, cwd=cwd, env=env, capture_output=True, text=True, timeout=timeout, shell=True)


def test_generated_project_runs_green(tmp_path, test_cases, saved_crawl, site_url):
    files = build_project(saved_crawl, test_cases, plan_tests(test_cases, saved_crawl, FakePlannerLLM()), "gen")
    for path, data in files.items():
        out = tmp_path / path
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_bytes(data)
    proj = tmp_path / "gen"

    # Turn TC-005 off in the input sheet to prove RunControl is honoured.
    xlsx = proj / "test-data" / "TestData.xlsx"
    wb = load_workbook(xlsx)
    for row in wb["RunControl"].iter_rows(min_row=2):
        if row[0].value == "TC-005":
            row[3].value = "N"
    wb.save(xlsx)

    r = _run("npm install --silent --no-audit --no-fund", proj)
    assert r.returncode == 0, r.stderr
    pin = os.getenv("PW_TEST_VERSION")
    if pin:
        assert _run(f"npm install --silent --no-audit --no-fund -D @playwright/test@{pin}", proj).returncode == 0

    env = {**os.environ, "BASE_URL": f"{site_url}/", "APP_USERNAME": "standard_user", "APP_PASSWORD": "secret_sauce"}
    r = _run("npx playwright test --reporter=line", proj, env)
    out = r.stdout + r.stderr
    assert r.returncode == 0, out[-3000:]
    assert "5 passed" in out and "1 skipped" in out
