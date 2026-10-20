import io
import json
import shutil
import subprocess

import pytest
from openpyxl import Workbook, load_workbook

from fakes import FakePlannerLLM
from pwgen.naming import element_name, page_class_name, url_template
from pwgen.planner import format_catalog, plan_tests, relevant_pages
from pwgen.renderer import build_project, zip_project
from pwgen.testcases import TestCaseFileError, read_test_cases, split_steps


# ---------------------------------------------------------------- naming

@pytest.mark.parametrize("url,key", [
    ("https://x.com/item/4?ref=a#top", "/item/:id"),
    ("https://x.com/item/5/", "/item/:id"),
    ("https://x.com/orders/3f2a9c1e-77aa-4bcd-9e10-1234567890ab", "/orders/:id"),
    ("https://x.com/", "/"),
    ("https://x.com/Inventory.html", "/inventory.html"),
])
def test_url_template_dedupes_ids(url, key):
    assert url_template(url) == key


def test_page_and_element_names():
    assert page_class_name("/checkout-step-one.html") == "CheckoutStepOnePage"
    assert page_class_name("/item/:id") == "ItemDetailPage"
    assert page_class_name("/") == "HomePage"
    taken = set()
    assert element_name("Login", "button", taken) == "loginButton"
    assert element_name("Login", "button", taken) == "loginButton2"
    assert element_name("shopping-cart-link", "button", set()) == "shoppingCartLink"
    assert element_name("delete", "text", set()) == "deleteText"
    assert element_name("page", "", set()) == "pageEl"  # would shadow BasePage.page


# ---------------------------------------------------------------- test case input

def test_reads_test_case_generator_export(test_cases):
    assert [t.id for t in test_cases] == ["TC-001", "TC-002", "TC-003", "TC-004", "TC-005", "TC-006"]
    assert test_cases[0].steps == ["Enter username", "Enter password", "Click the Login button"]
    assert test_cases[3].requirement_id == "FR-03"


def test_split_steps_variants():
    assert split_steps("1. Open\n2) Click\nStep 3: Check\n\n- keep") == ["Open", "Click", "Check", "- keep"]


def test_rejects_sheet_without_required_columns():
    wb = Workbook()
    wb.active.append(["Name", "Something"])
    wb.active.append(["a", "b"])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(TestCaseFileError, match="Missing column"):
        read_test_cases(buf.getvalue())


# ---------------------------------------------------------------- crawler (real browser, local site)

@pytest.fixture(scope="module")
def live_crawl(site_url):
    from pwgen.crawler import crawl

    return crawl(f"{site_url}/index.html", "standard_user", "secret_sauce")


def test_crawl_finds_unique_pages_and_logs_in(live_crawl):
    names = [p.name for p in live_crawl.pages]
    assert names == ["LoginPage", "InventoryPage", "CartPage", "OrdersPage", "ItemPage", "CheckoutStepOnePage"]
    assert live_crawl.login_page == "LoginPage"
    assert live_crawl.test_id_attribute == "data-test"


def test_crawl_dedupes_item_pages_and_stays_read_only(live_crawl):
    # Three product links lead to item.html?id=1..3: one page object.
    assert sum(p.url == "/item.html" for p in live_crawl.pages) == 1
    # Logout/reset are recorded as elements but never clicked: the crawl stayed logged in
    # long enough to reach checkout, which redirects to login when logged out.
    common = {e.name for e in live_crawl.common.elements}
    assert {"logoutSidebarLink", "resetSidebarLink", "shoppingCartLink", "shoppingCartBadgeText"} <= common
    log = "\n".join(live_crawl.log)
    assert "skip  button 'Logout': unsafe" in log and "'Add to cart': unsafe or changes data" in log


def test_crawl_follows_button_links_and_hidden_menus(live_crawl):
    log = "\n".join(live_crawl.log)
    # saucedemo-style cart: <button data-test="shopping-cart-link" aria-label="Cart">, no visible text.
    assert "click button 'Cart': -> /cart.html" in log
    # Orders is only in the hidden side menu: the crawler opens the menu first.
    assert "click button 'Orders': -> /orders.html" in log


def test_crawl_prefers_test_ids_and_shares_header(live_crawl):
    login = live_crawl.pages[0]
    by_name = {e.name: e for e in login.elements}
    assert (by_name["usernameInput"].strategy, by_name["usernameInput"].value) == ("testId", "username")
    inv = next(p for p in live_crawl.pages if p.name == "InventoryPage")
    assert "shoppingCartLink" not in {e.name for e in inv.elements}  # moved to CommonPage


def test_crawl_reports_bad_login(site_url):
    from pwgen.crawler import CrawlError, crawl

    with pytest.raises(CrawlError, match="Login didn't succeed"):
        crawl(f"{site_url}/index.html", "standard_user", "wrong")


# ---------------------------------------------------------------- planner

def test_planner_validates_against_crawl(test_cases, saved_crawl):
    llm = FakePlannerLLM()
    plans = {p.test_case_id: p for p in plan_tests(test_cases, saved_crawl, llm)}
    assert len(llm.calls) == 3  # 6 test cases, 2 per call
    tc6 = plans["TC-006"]
    assert [s.element for s in tc6.steps] == ["", "openMenuButton", "logoutSidebarLink", "loginButton"]
    assert any("noSuchElement" in u for u in tc6.unmapped)
    assert plans["TC-005"].unmapped == ["Check the product details (no specific value given)"]


def test_prompt_includes_relevant_pages_only(test_cases, saved_crawl):
    checkout_tc = test_cases[3]
    pages = [p.name for p in relevant_pages(saved_crawl, checkout_tc, n=2)]
    assert "CheckoutStepOnePage" in pages and "LoginPage" in pages
    catalog = format_catalog(saved_crawl, relevant_pages(saved_crawl, checkout_tc))
    assert catalog.startswith("CommonPage (on every page)")
    assert "(hidden until opened)" in catalog


# ---------------------------------------------------------------- renderer

@pytest.fixture
def project(test_cases, saved_crawl):
    plans = plan_tests(test_cases, saved_crawl, FakePlannerLLM())
    return build_project(saved_crawl, test_cases, plans, project_name="demo")


def test_project_layout(project):
    names = {p.split("/", 1)[1] for p in project}
    expected = {
        "package.json", "playwright.config.js", ".env.example", ".gitignore", "README.md",
        "pages/BasePage.js", "pages/CommonPage.js", "pages/LoginPage.js", "pages/InventoryPage.js",
        "object-repository/CommonPage.json", "object-repository/LoginPage.json",
        "fixtures/pages.js", "utils/testData.js", "test-data/TestData.xlsx",
        "tests/FR-01.spec.js", "tests/FR-05.spec.js",
    }
    assert expected <= names


@pytest.mark.skipif(not shutil.which("node"), reason="node not installed")
def test_every_js_file_is_valid(project, tmp_path):
    for path, data in project.items():
        if path.endswith(".js"):
            f = tmp_path / path.replace("/", "_")
            f.write_bytes(data)
            r = subprocess.run(["node", "--check", str(f)], capture_output=True, text=True)
            assert r.returncode == 0, f"{path}: {r.stderr}"


def test_object_repository_and_input_sheet(project):
    repo = json.loads(project["demo/object-repository/LoginPage.json"])
    assert repo["url"] == "/index.html"
    assert repo["elements"]["usernameInput"] == {
        "strategy": "testId", "value": "username", "kind": "input", "description": 'Input "Username"'}
    wb = load_workbook(io.BytesIO(project["demo/test-data/TestData.xlsx"]))
    assert wb.sheetnames == ["RunControl", "TestData"]
    rows = list(wb["TestData"].iter_rows(min_row=2, values_only=True))
    assert ("TC-004", "postalCode", "637211") in rows
    assert all(r[3] == "Y" for r in wb["RunControl"].iter_rows(min_row=2, values_only=True))


def test_specs_use_page_objects_and_mark_todos(project):
    spec = project["demo/tests/FR-03.spec.js"].decode()
    assert "await loginPage.login();" in spec
    assert "await checkoutStepOnePage.lastNameInput.fill(data.lastName);" in spec
    assert "await expect(checkoutStepOnePage.errorText).toContainText(data.expectedError);" in spec
    fr5 = project["demo/tests/FR-05.spec.js"].decode()
    assert "// TODO (not automated)" in fr5
    # Incomplete test: automated steps first, then fixme so it can't show as passed.
    assert fr5.rstrip().endswith(
        "test.fixme(true, \"1 step(s) not automated yet: complete the TODOs above\");\n  });\n});")
    assert "test.fixme" not in project["demo/tests/FR-03.spec.js"].decode()  # complete test
    # A click that leads to another page is followed by a wait for that page, so the next
    # assertion can't hit the old page while it unloads (strict-mode failure on saucedemo).
    detail = project["demo/tests/FR-04.spec.js"].decode()
    assert ("await inventoryPage.trailBackpackLink.click();\n"
            "    // Product name is shown\n"
            "    await itemPage.waitForPage();\n"
            "    await expect(itemPage.inventoryItemNameText).toBeVisible();") in detail
    tc1 = project["demo/tests/FR-01.spec.js"].decode().split("TC-002")[0]
    assert tc1.count("waitForPage") == 1  # after Login -> inventory; not before the first goto
    assert "await loginPage.waitForPage" not in tc1
    login = project["demo/pages/LoginPage.js"].decode()
    assert "async login(username = process.env.APP_USERNAME" in login
    assert "secret_sauce" not in project["demo/.env.example"].decode()


def test_zip_round_trip(project):
    import zipfile

    z = zipfile.ZipFile(io.BytesIO(zip_project(project)))
    assert sorted(z.namelist()) == sorted(project)
