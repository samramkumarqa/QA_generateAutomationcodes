"""Read test cases from the AI Test Case Generator's Excel test pack (or a sheet shaped like it)."""

import io
import re

from openpyxl import load_workbook

from .schemas import TestCaseIn


class TestCaseFileError(ValueError):
    __test__ = False  # not a pytest test class


# Header names we accept for each field (lower-case, spaces collapsed).
COLUMNS = {
    "id": ["test case id", "tc id", "id", "test id"],
    "requirement_id": ["requirement id", "req id", "requirement"],
    "title": ["title", "test case title", "test case", "summary", "scenario"],
    "preconditions": ["preconditions", "pre-conditions", "precondition"],
    "steps": ["steps", "test steps", "step"],
    "test_data": ["test data", "data"],
    "expected_result": ["expected result", "expected results", "expected"],
}


def _norm(h) -> str:
    return re.sub(r"\s+", " ", str(h or "")).strip().lower()


def split_steps(text: str) -> list[str]:
    lines = [ln.strip() for ln in str(text or "").replace("\r", "").split("\n")]
    steps = [re.sub(r"^\s*(step\s*)?\d+\s*[.):-]\s*", "", ln, flags=re.I).strip() for ln in lines]
    return [s for s in steps if s]


def read_test_cases(data: bytes) -> list[TestCaseIn]:
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise TestCaseFileError(f"Could not open the Excel file: {exc}") from exc

    ws = wb["Test Cases"] if "Test Cases" in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise TestCaseFileError("The sheet is empty.")
    header = [_norm(h) for h in rows[0]]
    col = {}
    for field, names in COLUMNS.items():
        for i, h in enumerate(header):
            if h in names:
                col[field] = i
                break
    missing = [f for f in ("id", "title", "steps") if f not in col]
    if missing:
        raise TestCaseFileError(
            "Missing column(s): " + ", ".join(missing).replace("_", " ")
            + ". Use the Excel exported by the AI Test Case Generator, or a sheet with "
            "'Test Case ID', 'Title' and 'Steps' columns."
        )

    out: list[TestCaseIn] = []
    for row in rows[1:]:
        get = lambda f: str(row[col[f]] or "").strip() if f in col and col[f] < len(row) else ""
        if not get("id") and not get("title"):
            continue
        steps = split_steps(get("steps"))
        if not steps:
            continue
        out.append(
            TestCaseIn(
                id=get("id") or f"TC-{len(out) + 1:03d}",
                requirement_id=get("requirement_id"),
                title=get("title"),
                preconditions=get("preconditions"),
                steps=steps,
                test_data=get("test_data"),
                expected_result=get("expected_result"),
            )
        )
    if not out:
        raise TestCaseFileError("No test cases with steps were found in the sheet.")
    return out
