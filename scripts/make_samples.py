"""Build samples/sample_testcases.xlsx in the AI Test Case Generator's export format.

The test cases target Swag Labs (saucedemo.com), the public demo site, and also run against
the local practice shop in tests/site used by the test suite.

Run:  python scripts/make_samples.py
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

OUT = Path(__file__).resolve().parent.parent / "samples" / "sample_testcases.xlsx"

HEADERS = [
    "Test Case ID", "Requirement ID", "Title", "Type", "Priority", "Preconditions", "Steps",
    "Test Data", "Expected Result", "Status", "Actual Result", "Revision",
]

CASES = [
    ("TC-001", "FR-01", "Verify successful login with valid credentials", "Functional", "High",
     "User is on the login page",
     ["Enter username", "Enter password", "Click the Login button"],
     "username: standard_user; password: secret_sauce",
     "The Products page is displayed"),
    ("TC-002", "FR-01", "Verify a locked out user cannot log in", "Negative", "High",
     "User is on the login page",
     ["Enter username of the locked out user", "Enter password", "Click the Login button"],
     "username: locked_out_user; password: secret_sauce",
     "Error 'Epic sadface: Sorry, this user has been locked out.' is shown"),
    ("TC-003", "FR-02", "Verify adding the backpack to the cart updates the cart badge", "Functional", "High",
     "User is logged in and the cart is empty",
     ["On the Products page, click Add to cart for the backpack", "Check the cart icon in the header"],
     "",
     "The cart badge shows 1"),
    ("TC-004", "FR-03", "Verify checkout requires the first name", "Negative", "High",
     "User is logged in",
     ["Add the backpack to the cart", "Open the cart", "Click Checkout",
      "Enter last name and postal code, leave first name empty", "Click Continue"],
     "lastName: Kumar; postalCode: 637211",
     "Error 'Error: First Name is required' is shown"),
    ("TC-005", "FR-04", "Verify the product details page opens from the product list", "Functional", "Medium",
     "User is logged in",
     ["Click the backpack product name", "Check the product details", "Click Back to products"],
     "",
     "Product name is shown on the details page; Back to products returns to the Products page"),
    ("TC-006", "FR-05", "Verify the user can log out", "Functional", "Medium",
     "User is logged in",
     ["Open the menu", "Click Logout"],
     "",
     "The login page is shown"),
]


def main() -> None:
    OUT.parent.mkdir(exist_ok=True)
    wb = Workbook()
    wb.active.title = "Summary"
    wb.active["A1"] = "Sample approved test pack (AI Test Case Generator format)"
    ws = wb.create_sheet("Test Cases")
    ws.append(HEADERS)
    for tc_id, req, title, typ, prio, pre, steps, data, exp in CASES:
        ws.append([tc_id, req, title, typ, prio, pre,
                   "\n".join(f"{i}. {s}" for i, s in enumerate(steps, 1)), data, exp, "Not Run", "", 1])
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3A5F")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDEFGHIJKL", [12, 12, 45, 12, 9, 30, 55, 35, 45, 9, 12, 9]):
        ws.column_dimensions[col].width = w
    wb.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
